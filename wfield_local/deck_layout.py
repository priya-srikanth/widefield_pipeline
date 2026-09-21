"""THE DECK'S DRAWING SURFACE — one slide canvas that owns the presentation and the bookkeeping.

Split out of `locanmf_analysis_deck` on 2026-09-21. These were eleven closures inside
`build_analysis_deck`, sharing eleven mutable locals between them: the presentation, the blank
layout, the slide dimensions, the placed/missing counters, the figure manifest rows, the slide
order, the per-slide figure index, the methods-dedup table and the queue of pending notes. A
closure over eleven variables is a class that has not admitted it, and it could not be tested or
reused while it lived inside a 2,000-line function.

WHAT THIS CLASS IS AND IS NOT. It knows how to put a thing on a slide and it counts what it put
there. It does NOT know what figures exist, what order the sections go in, which sessions were
recorded, or when to refuse to publish -- all of that stays in the builder, which is the part
that is actually about this study.

THE BOOKKEEPING IS THE POINT, not a side effect. `placed` / `missing_figures` feed the
completeness gate that refuses to overwrite a published deck with a worse one, and
`placed_figures` feeds the manifest that is how an ORPHANED reference surfaces -- a slide pointing
at a filename no step writes any more, invisible to every other check because the file is present,
just never updated. Any new way of putting a figure on a slide must go through `exists()` or it
will be silently exempt from both.

NOTES ARE QUEUED, NOT WRITTEN, and that is load-bearing -- see `note`.

THE BUILDER BINDS THESE AS LOCAL NAMES (`slide, title, note, big, ... = canvas.slide, ...`) rather
than writing `canvas.title(...)` at each of ~250 call sites. That keeps the section code reading
as the narrative it is -- `title(s, ...)`, `big(s, fig)` -- instead of putting the same receiver in
front of every line of it, and it meant the split moved these definitions without touching the
body, so the 531-slide fingerprint that proved the deck unchanged was a real test of the move
rather than a test of a rewrite.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from wfield_local import deck_values

NAVY = RGBColor(0x1F, 0x33, 0x55)
GREY = RGBColor(0x55, 0x55, 0x55)
SLATE = RGBColor(0x44, 0x55, 0x77)


class SlideCanvas:
    """A 13.333 x 7.5 in presentation plus everything the deck counts while filling it."""

    def __init__(self, values: deck_values.Resolver):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.BLANK = self.prs.slide_layouts[6]
        self.SW, self.SH = self.prs.slide_width, self.prs.slide_height
        self.placed = {"present": 0, "missing": 0}
        self.missing_figures = []
        self.placed_figures = []
        self.slide_order = []
        self.figs_by_slide = {}
        self.seen_methods = {}
        self.pending_notes = []
        #: Sidecar resolver for `note`; the BUILDER chooses its search roots, because the
        #: precedence (grant sets, then epoch, then the working dir) is the same precedence the
        #: placement patterns use and belongs with them.
        self.values = values

    # ------------------------------------------------------------------ slides and bookkeeping

    def slide(self):
        sl = self.prs.slides.add_slide(self.BLANK)
        self.slide_order.append(sl)
        return sl

    def _record(self, sl, p):
        self.figs_by_slide.setdefault(id(sl._element), []).append(Path(p).name)

    def exists(self, p):
        """Does this figure exist -- and count it either way, present or missing.

        Renamed from `_exists` in the 2026-09-21 move: the builder calls it directly as well as
        `big` and `grid` doing so, and a name that five call sites outside the class already used
        was never private.
        """
        ok = Path(p).exists()
        self.placed["present" if ok else "missing"] += 1
        if not ok:
            self.missing_figures.append(Path(p).name)   # NAME the gap: a count alone cannot be acted on
        else:
            self.placed_figures.append((Path(p).name, Path(p).stat().st_mtime))
        return ok

    # ----------------------------------------------------------------------------------- text

    def title(self, s, text, sub=None, trials=None):
        """Slide title, an optional subtitle, and an optional TRIAL POPULATION line.

        The population is its own line rather than a clause in the subtitle because it is the one
        fact a reader needs before comparing two slides, and subtitles here are already long enough
        that it would be buried in the middle of one.
        """
        tf = s.shapes.add_textbox(Inches(0.4), Inches(0.16), Inches(12.6), Inches(1.15)).text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if sub:
            r2 = tf.add_paragraph().add_run()
            r2.text = sub
            r2.font.size = Pt(12.5)
            r2.font.color.rgb = GREY
        if trials:
            r3 = tf.add_paragraph().add_run()
            r3.text = trials
            r3.font.size = Pt(10)
            r3.font.italic = True
            r3.font.color.rgb = SLATE

    def bullets(self, s, items, top=1.5, size=13.5, width=12.4):
        """A text-only slide body. Sections A-F are all figures, but the post-stroke section has to
        state what is comparable to what BEFORE showing a number -- that argument has no figure, and
        burying it in the speaker notes is how the first version of this analysis shipped a headline
        that was mostly trial composition."""
        tf = s.shapes.add_textbox(Inches(0.45), Inches(top), Inches(width),
                                  self.SH - Inches(top) - Inches(0.3)).text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = para.add_run()
            r.text = "•  " + it
            r.font.size = Pt(size)
            para.space_after = Pt(9)

    def divider(self, text, sub=None):
        s = self.slide()
        tf = s.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.9)).text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if sub:
            r2 = tf.add_paragraph().add_run()
            r2.text = sub
            r2.font.size = Pt(15)
            r2.font.color.rgb = GREY

    # -------------------------------------------------------------------------------- figures

    def big(self, s, p, top=1.4, width=12.7, bottom=0.15):
        """Place one figure, scaled to fit the slide in BOTH dimensions.

        Scaling by width alone overflows the bottom of the slide whenever a figure is taller than
        (13.333 - margins) : (7.5 - top), which is most multi-row figures -- the picture simply ran
        off the deck and the axis labels at the foot of it were never visible. Height is now capped
        at the space actually available and the width follows from the image's own aspect ratio, so
        a figure is never cropped and its fonts shrink proportionally rather than disappearing.
        """
        # THE `_record` CALL USED TO SIT ABOVE THIS DOCSTRING, which made the docstring a bare
        # string expression and left `big.__doc__` as None -- the explanation above was invisible
        # to `help()` and to every doc tool. Moving the call below it changes nothing at runtime.
        self._record(s, p)
        if not self.exists(p):
            return
        with Image.open(str(p)) as im:
            iw, ih = im.size
        avail_h = float(self.SH.inches) - top - bottom
        w_in = min(float(width), avail_h * (iw / ih))
        w = Inches(w_in)
        s.shapes.add_picture(str(p), (self.SW - w) / 2, Inches(top), width=w)

    def grid(self, s, paths, cols=2, top=1.25, side=0.25, gap=0.18, bottom=0.25):
        paths = [Path(p) for p in paths]
        for _p in paths:
            self._record(s, _p)
        present = [p for p in paths if self.exists(p)]
        if not present:
            return
        rows = (len(present) + cols - 1) // cols
        cell_w = (self.SW - Inches(side) * 2 - Inches(gap) * (cols - 1)) / cols
        cell_h = (self.SH - Inches(top) - Inches(bottom) - Inches(gap) * (rows - 1)) / rows
        for i, p in enumerate(present):
            r, c = divmod(i, cols)
            iw, ih = Image.open(str(p)).size
            scale = min(cell_w / iw, cell_h / ih)
            w, h = int(iw * scale), int(ih * scale)
            left = Inches(side) + c * (cell_w + Inches(gap)) + (cell_w - w) / 2
            t = Inches(top) + r * (cell_h + Inches(gap)) + (cell_h - h) / 2
            s.shapes.add_picture(str(p), left, t, width=w, height=h)

    # --------------------------------------------------------------------------- speaker notes

    def note(self, s, text, specific=None):
        """Queue this slide's speaker notes. Written for real by `flush_notes` before the save.

        DEFERRED, AND IT HAS TO BE. A note may quote its own figure's sidecar via a ``SELF`` token
        (see `deck_values`), and many slides call `note` BEFORE placing the figure -- on the first
        real build that left `[[? SELF has no figure on this slide]]` on the two converted notes,
        because `figs_by_slide` was still empty for that slide. Reordering every call site would fix
        it only until the next one was written in the old order.

        RESOLUTION CANNOT SIMPLY MOVE AFTER THE DEDUP EITHER. The methods block is deduped by
        hashing its text, and two arms whose prose is identical differ only in their resolved
        numbers; hashing the RAW text would collapse them and replace the second with "same as slide
        N", pointing the reader at another arm's numbers. So resolve first, then hash -- which means
        both must wait until every figure is recorded.
        """
        self.pending_notes.append((s, len(self.prs.slides), text, specific))

    def flush_notes(self):
        """Resolve every queued note against its slide's figure, dedupe, and write."""
        for s, idx, text, specific in self.pending_notes:
            self._write_note(s, idx, text, specific)

    def _write_note(self, s, idx, text, specific):
        # `SELF` in a token means THIS SLIDE'S figure. Most notes are placed by a glob and serve
        # every trial-class arm, so a hard-coded stem would print one arm's numbers onto all of them.
        _self = deck_values.stem_of((self.figs_by_slide.get(id(s._element)) or [None])[0])
        text = self.values.resolve(text, self_stem=_self)
        specific = self.values.resolve(specific, self_stem=_self)
        parts = []
        if specific:
            parts.append("THIS SLIDE" + chr(10) + specific.strip())
        # HASH THE WHOLE TEXT, not a prefix. This keyed on text[:80] until 2026-08-23, when
        # _M_LICK_UNIT was PREPENDED to M_FIXEDSCALE, M_GATE and M_POSTSTROKE -- three unrelated
        # methods blocks that then shared their first 80 characters. The dedup would have called the
        # second and third "same as slide N" and pointed each at the FIRST one's methods: a wrong
        # cross-reference reads exactly like a right one, which is worse than the repetition this
        # replaced.
        key = hashlib.sha1((text or "").encode("utf-8")).hexdigest()
        if not text:
            pass
        elif key in self.seen_methods:
            parts.append(f"METHODS -- same as slide {self.seen_methods[key]}; "
                         f"not repeated here.")
        else:
            self.seen_methods[key] = idx
            parts.append("METHODS" + chr(10) + text.strip())
        s.notes_slide.notes_text_frame.text = (chr(10) + chr(10)).join(parts)
